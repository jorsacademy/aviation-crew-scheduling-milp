"""Synthetic crew scheduling MILP example using PuLP.

This module modernizes an older teaching example. It is intended for operations-
research education and demonstration, not operational deployment.
"""

from __future__ import annotations

import copy
import csv
import time
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
import pulp as pl


class CrewSchedulingOptimizer:
    """Mixed-integer crew scheduling model with qualifications, duty and rest rules."""

    def __init__(self, *, enforce_rest: bool = True, min_high_completion_ratio: float = 0.60):
        self.enforce_rest = enforce_rest
        self.min_high_completion_ratio = min_high_completion_ratio
        self.setup_data()
        self.create_model()

    def setup_data(self) -> None:
        self.missions: dict[str, dict[str, Any]] = {
            "M1": {"day": 1, "duration": 14, "aircraft": 2, "priority": "High", "combat": True, "night": False},
            "M2": {"day": 1, "duration": 8, "aircraft": 1, "priority": "Medium", "combat": False, "night": False},
            "M3": {"day": 2, "duration": 16, "aircraft": 1, "priority": "High", "combat": False, "night": True},
            "M4": {"day": 2, "duration": 10, "aircraft": 2, "priority": "Medium", "combat": False, "night": False},
            "M5": {"day": 3, "duration": 12, "aircraft": 1, "priority": "Critical", "combat": True, "night": True},
            "M6": {"day": 4, "duration": 6, "aircraft": 1, "priority": "Low", "combat": False, "night": False},
            "M7": {"day": 4, "duration": 18, "aircraft": 2, "priority": "High", "combat": True, "night": False},
            "M8": {"day": 5, "duration": 8, "aircraft": 1, "priority": "Medium", "combat": False, "night": True},
            "M9": {"day": 5, "duration": 14, "aircraft": 1, "priority": "High", "combat": False, "night": False},
            "M10": {"day": 6, "duration": 10, "aircraft": 2, "priority": "Medium", "combat": False, "night": False},
            "M11": {"day": 7, "duration": 12, "aircraft": 1, "priority": "High", "combat": True, "night": False},
            "M12": {"day": 7, "duration": 8, "aircraft": 1, "priority": "Low", "combat": False, "night": False},
        }

        self.pilots = [f"P{i}" for i in range(1, 9)]
        self.copilots = [f"C{i}" for i in range(1, 9)]
        self.engineers = [f"E{i}" for i in range(1, 5)]
        self.crew_members = self.pilots + self.copilots + self.engineers

        self.qualifications = {
            "P1": {"combat": True, "night": True}, "P2": {"combat": True, "night": True},
            "P3": {"combat": True, "night": True}, "P4": {"combat": True, "night": False},
            "P5": {"combat": True, "night": False}, "P6": {"combat": False, "night": True},
            "P7": {"combat": False, "night": True}, "P8": {"combat": False, "night": True},
            "C1": {"combat": True, "night": True}, "C2": {"combat": True, "night": True},
            "C3": {"combat": True, "night": False}, "C4": {"combat": True, "night": False},
            "C5": {"combat": False, "night": True}, "C6": {"combat": False, "night": True},
            "C7": {"combat": False, "night": True}, "C8": {"combat": False, "night": False},
            "E1": {"combat": True, "night": True}, "E2": {"combat": True, "night": False},
            "E3": {"combat": False, "night": True}, "E4": {"combat": False, "night": True},
        }
        self.unavailable = {"P3": [2, 3], "C6": [5, 6], "E2": [4]}

        self.base_cost = 200.0
        self.overtime_multiplier = 1.5
        self.hazard_pay = 100.0
        self.penalties = {"Critical": 50000.0, "High": 25000.0, "Medium": 10000.0, "Low": 2500.0}
        self.days = list(range(1, 8))
        self.max_daily_hours = 16
        self.max_weekly_hours = 60
        self.regular_daily_hours = 8

    def create_model(self) -> None:
        self.model = pl.LpProblem("Crew_Scheduling", pl.LpMinimize)
        self.create_variables()
        self.create_objective()
        self.create_constraints()

    def create_variables(self) -> None:
        self.x = {
            (m, a, c): pl.LpVariable(f"x_{m}_{a}_{c}", cat="Binary")
            for m, info in self.missions.items()
            for a in range(1, info["aircraft"] + 1)
            for c in self.crew_members
        }
        self.y = {m: pl.LpVariable(f"y_{m}", cat="Binary") for m in self.missions}
        self.regular_hours = {}
        self.overtime_hours = {}
        self.total_hours = {}
        self.worked = {}
        self.combat_worked = {}
        for c in self.crew_members:
            for d in self.days:
                self.regular_hours[c, d] = pl.LpVariable(f"regular_{c}_{d}", 0, self.regular_daily_hours)
                self.overtime_hours[c, d] = pl.LpVariable(f"overtime_{c}_{d}", 0)
                self.total_hours[c, d] = pl.LpVariable(f"total_{c}_{d}", 0)
                self.worked[c, d] = pl.LpVariable(f"worked_{c}_{d}", cat="Binary")
                self.combat_worked[c, d] = pl.LpVariable(f"combat_worked_{c}_{d}", cat="Binary")

    def create_objective(self) -> None:
        labor = pl.lpSum(
            self.base_cost * self.regular_hours[c, d]
            + self.base_cost * self.overtime_multiplier * self.overtime_hours[c, d]
            for c in self.crew_members for d in self.days
        )
        hazard = pl.lpSum(
            self.hazard_pay * self.missions[m]["duration"] * self.x[m, a, c]
            for m, info in self.missions.items() if info["combat"]
            for a in range(1, info["aircraft"] + 1)
            for c in self.crew_members
        )
        penalties = pl.lpSum(
            self.penalties[info["priority"]] * (1 - self.y[m])
            for m, info in self.missions.items()
        )
        self.model += labor + hazard + penalties

    def create_constraints(self) -> None:
        for m, info in self.missions.items():
            for a in range(1, info["aircraft"] + 1):
                self.model += pl.lpSum(self.x[m, a, p] for p in self.pilots) == self.y[m], f"pilot_{m}_{a}"
                self.model += pl.lpSum(self.x[m, a, c] for c in self.copilots) == self.y[m], f"copilot_{m}_{a}"

                engineer_required = info["duration"] > 10 or info["combat"]
                eng_sum = pl.lpSum(self.x[m, a, e] for e in self.engineers)
                self.model += eng_sum == self.y[m] if engineer_required else eng_sum == 0, f"engineer_{m}_{a}"

                for c in self.crew_members:
                    if info["combat"] and not self.qualifications[c]["combat"]:
                        self.model += self.x[m, a, c] == 0
                    if info["night"] and not self.qualifications[c]["night"]:
                        self.model += self.x[m, a, c] == 0
                    if info["day"] in self.unavailable.get(c, []):
                        self.model += self.x[m, a, c] == 0

            for c in self.crew_members:
                self.model += pl.lpSum(self.x[m, a, c] for a in range(1, info["aircraft"] + 1)) <= 1, f"single_aircraft_{m}_{c}"

        # Without mission start/end times, assume a crew member can serve at most one mission per day.
        for c in self.crew_members:
            for d in self.days:
                missions_today = [m for m, info in self.missions.items() if info["day"] == d]
                assignment_terms = [
                    self.x[m, a, c]
                    for m in missions_today
                    for a in range(1, self.missions[m]["aircraft"] + 1)
                ]
                if assignment_terms:
                    self.model += pl.lpSum(assignment_terms) <= 1, f"one_mission_per_day_{c}_{d}"

                self.model += self.total_hours[c, d] == pl.lpSum(
                    self.missions[m]["duration"] * self.x[m, a, c]
                    for m in missions_today
                    for a in range(1, self.missions[m]["aircraft"] + 1)
                ), f"total_hours_{c}_{d}"
                self.model += self.regular_hours[c, d] + self.overtime_hours[c, d] == self.total_hours[c, d], f"time_split_{c}_{d}"
                self.model += self.total_hours[c, d] <= self.max_daily_hours, f"daily_limit_{c}_{d}"
                self.model += self.total_hours[c, d] <= self.max_daily_hours * self.worked[c, d]
                self.model += self.total_hours[c, d] >= self.worked[c, d]

                combat_terms = [
                    self.x[m, a, c]
                    for m in missions_today if self.missions[m]["combat"]
                    for a in range(1, self.missions[m]["aircraft"] + 1)
                ]
                if combat_terms:
                    for term in combat_terms:
                        self.model += self.combat_worked[c, d] >= term
                    self.model += self.combat_worked[c, d] <= pl.lpSum(combat_terms)
                else:
                    self.model += self.combat_worked[c, d] == 0

        for c in self.crew_members:
            self.model += pl.lpSum(self.total_hours[c, d] for d in self.days) <= self.max_weekly_hours, f"weekly_limit_{c}"

        if self.enforce_rest:
            for c in self.crew_members:
                for d in range(1, 7):
                    self.model += self.total_hours[c, d] <= 8 + 8 * (1 - self.worked[c, d + 1]), f"rest_{c}_{d}"
                    self.model += self.combat_worked[c, d] + self.worked[c, d + 1] <= 1, f"combat_rest_{c}_{d}"

        for m, info in self.missions.items():
            if info["priority"] == "Critical":
                self.model += self.y[m] == 1, f"force_critical_{m}"

        high = [m for m, info in self.missions.items() if info["priority"] == "High"]
        minimum = max(1, int(np.ceil(self.min_high_completion_ratio * len(high))))
        self.model += pl.lpSum(self.y[m] for m in high) >= minimum, "minimum_high_priority_completion"

    def solve(self, solver: pl.LpSolver | None = None, time_limit: int = 300, msg: bool = False) -> int:
        if solver is not None:
            return self.model.solve(solver)

        candidates = [
            ("GUROBI", lambda: pl.GUROBI_CMD(timeLimit=time_limit, msg=msg)),
            ("CPLEX", lambda: pl.CPLEX_CMD(timeLimit=time_limit, msg=msg)),
            ("CBC", lambda: pl.PULP_CBC_CMD(timeLimit=time_limit, msg=msg)),
        ]
        errors: list[str] = []
        for name, factory in candidates:
            try:
                candidate = factory()
                if candidate.available():
                    return self.model.solve(candidate)
            except Exception as exc:  # solver-specific availability/runtime failure
                errors.append(f"{name}: {exc}")

        detail = "; ".join(errors) if errors else "no configured solver executable was available"
        raise RuntimeError(f"Unable to solve model: {detail}")

    def get_results(self) -> dict[str, Any]:
        status_name = pl.LpStatus[self.model.status]
        if self.model.status != pl.LpStatusOptimal:
            raise RuntimeError(f"No optimal solution available (status={status_name}).")

        mission_assignments: dict[str, dict[str, list[str]]] = {}
        completion: dict[str, int] = {}
        crew_schedules: dict[str, dict[str, dict[str, float]]] = {}

        for m, info in self.missions.items():
            completion[m] = int(round(pl.value(self.y[m])))
            mission_assignments[m] = {}
            for a in range(1, info["aircraft"] + 1):
                mission_assignments[m][f"Aircraft_{a}"] = [
                    c for c in self.crew_members if pl.value(self.x[m, a, c]) > 0.5
                ]

        for c in self.crew_members:
            schedule = {}
            for d in self.days:
                total = float(pl.value(self.total_hours[c, d]) or 0.0)
                if total > 1e-6:
                    schedule[f"Day_{d}"] = {
                        "total_hours": total,
                        "regular_hours": float(pl.value(self.regular_hours[c, d]) or 0.0),
                        "overtime_hours": float(pl.value(self.overtime_hours[c, d]) or 0.0),
                    }
            crew_schedules[c] = schedule

        labor = sum(
            self.base_cost * float(pl.value(self.regular_hours[c, d]) or 0.0)
            + self.base_cost * self.overtime_multiplier * float(pl.value(self.overtime_hours[c, d]) or 0.0)
            for c in self.crew_members for d in self.days
        )
        hazard = sum(
            self.hazard_pay * info["duration"] * float(pl.value(self.x[m, a, c]) or 0.0)
            for m, info in self.missions.items() if info["combat"]
            for a in range(1, info["aircraft"] + 1)
            for c in self.crew_members
        )
        penalty = sum(self.penalties[info["priority"]] * (1 - completion[m]) for m, info in self.missions.items())

        return {
            "status": status_name,
            "objective_value": float(pl.value(self.model.objective)),
            "mission_assignments": mission_assignments,
            "crew_schedules": crew_schedules,
            "mission_completion": completion,
            "cost_breakdown": {"labor": labor, "hazard": hazard, "penalties": penalty},
        }

    def clone_for_scenario(self) -> "CrewSchedulingOptimizer":
        clone = copy.deepcopy(self)
        clone.create_model()
        return clone

    def sensitivity_analysis(self, time_limit: int = 120) -> list[dict[str, Any]]:
        if self.model.status != pl.LpStatusOptimal:
            raise RuntimeError("Solve the base model before sensitivity analysis.")
        base = float(pl.value(self.model.objective))
        outputs = []

        scenarios = []
        reduced = CrewSchedulingOptimizer(enforce_rest=self.enforce_rest, min_high_completion_ratio=self.min_high_completion_ratio)
        reduced.pilots = reduced.pilots[:-2]
        reduced.crew_members = reduced.pilots + reduced.copilots + reduced.engineers
        reduced.create_model()
        scenarios.append(("Reduced crew (remove 2 pilots)", reduced))

        penalty = CrewSchedulingOptimizer(enforce_rest=self.enforce_rest, min_high_completion_ratio=self.min_high_completion_ratio)
        penalty.penalties = {k: 2 * v for k, v in penalty.penalties.items()}
        penalty.create_model()
        scenarios.append(("Double mission penalties", penalty))

        relaxed = CrewSchedulingOptimizer(enforce_rest=False, min_high_completion_ratio=self.min_high_completion_ratio)
        scenarios.append(("Relaxed rest constraints", relaxed))

        for name, scenario in scenarios:
            try:
                status = scenario.solve(time_limit=time_limit, msg=False)
                status_name = pl.LpStatus[status]
                value = float(pl.value(scenario.model.objective)) if status == pl.LpStatusOptimal else None
                outputs.append({
                    "scenario": name,
                    "status": status_name,
                    "objective": value,
                    "change_pct": ((value - base) / base * 100) if value is not None and base else None,
                })
            except Exception as exc:
                outputs.append({"scenario": name, "status": "ERROR", "objective": None, "change_pct": None, "error": str(exc)})
        return outputs

    def print_solution(self) -> None:
        r = self.get_results()
        print("=" * 64)
        print("CREW SCHEDULING SOLUTION")
        print("=" * 64)
        print(f"Objective: ${r['objective_value']:,.2f}")
        print("Cost breakdown:", ", ".join(f"{k}=${v:,.2f}" for k, v in r["cost_breakdown"].items()))
        for m, aircraft_map in r["mission_assignments"].items():
            info = self.missions[m]
            marker = "✓" if r["mission_completion"][m] else "✗"
            print(f"\n{m} {marker} | day={info['day']} duration={info['duration']}h priority={info['priority']}")
            for aircraft, crew in aircraft_map.items():
                if crew:
                    print(f"  {aircraft}: {', '.join(crew)}")


class AdvancedAnalysis:
    def __init__(self, optimizer: CrewSchedulingOptimizer):
        self.optimizer = optimizer
        self.results = optimizer.get_results()

    def workload_distribution(self) -> dict[str, Any]:
        workloads = {
            c: sum(day["total_hours"] for day in self.results["crew_schedules"][c].values())
            for c in self.optimizer.crew_members
        }
        values = list(workloads.values())
        avg = float(np.mean(values)) if values else 0.0
        std = float(np.std(values)) if values else 0.0
        return {"by_crew": workloads, "average": avg, "std_dev": std, "coefficient_of_variation": std / avg if avg else 0.0}

    def mission_completion_by_priority(self) -> dict[str, dict[str, float]]:
        stats: dict[str, dict[str, float]] = defaultdict(lambda: {"completed": 0, "total": 0, "rate": 0.0})
        for m, completed in self.results["mission_completion"].items():
            p = self.optimizer.missions[m]["priority"]
            stats[p]["total"] += 1
            stats[p]["completed"] += completed
        for values in stats.values():
            values["rate"] = 100.0 * values["completed"] / values["total"] if values["total"] else 0.0
        return dict(stats)


def export_results_to_csv(optimizer: CrewSchedulingOptimizer, filename: str = "crew_schedule_results.csv") -> Path:
    results = optimizer.get_results()
    path = Path(filename)
    rows = []
    for m, assignment in results["mission_assignments"].items():
        info = optimizer.missions[m]
        for aircraft, crew_list in assignment.items():
            for crew in crew_list:
                rows.append({
                    "Mission": m, "Aircraft": aircraft, "Crew_Member": crew,
                    "Day": info["day"], "Duration": info["duration"], "Priority": info["priority"],
                    "Combat": info["combat"], "Night": info["night"], "Completed": results["mission_completion"][m],
                })
    fieldnames = ["Mission", "Aircraft", "Crew_Member", "Day", "Duration", "Priority", "Combat", "Night", "Completed"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return path


def export_results_to_excel(optimizer: CrewSchedulingOptimizer, filename: str = "crew_schedule_results.xlsx") -> Path:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl. Install it with: pip install openpyxl") from exc

    results = optimizer.get_results()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mission Assignments"
    ws.append(["Mission", "Aircraft", "Crew Member", "Day", "Duration", "Priority", "Completed"])
    for m, assignment in results["mission_assignments"].items():
        info = optimizer.missions[m]
        for aircraft, crew_list in assignment.items():
            if crew_list:
                for crew in crew_list:
                    ws.append([m, aircraft, crew, info["day"], info["duration"], info["priority"], results["mission_completion"][m]])
            else:
                ws.append([m, aircraft, "", info["day"], info["duration"], info["priority"], results["mission_completion"][m]])

    ws_crew = wb.create_sheet("Crew Schedules")
    ws_crew.append(["Crew Member", "Day", "Total Hours", "Regular Hours", "Overtime Hours"])
    for crew, schedule in results["crew_schedules"].items():
        for day, values in schedule.items():
            ws_crew.append([crew, day, values["total_hours"], values["regular_hours"], values["overtime_hours"]])

    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Status", results["status"]])
    ws_summary.append(["Objective", results["objective_value"]])
    for key, value in results["cost_breakdown"].items():
        ws_summary.append([f"Cost - {key}", value])

    path = Path(filename)
    wb.save(path)
    return path


def performance_test(runs: int = 3, time_limit: int = 60) -> dict[str, Any]:
    times = []
    statuses = []
    for _ in range(runs):
        start = time.perf_counter()
        optimizer = CrewSchedulingOptimizer()
        status = optimizer.solve(time_limit=time_limit, msg=False)
        times.append(time.perf_counter() - start)
        statuses.append(pl.LpStatus[status])
    return {"runs": runs, "statuses": statuses, "average_seconds": float(np.mean(times)), "times": times}


def main() -> None:
    optimizer = CrewSchedulingOptimizer()
    print(f"Variables: {len(optimizer.model.variables())}")
    print(f"Constraints: {len(optimizer.model.constraints)}")
    status = optimizer.solve(time_limit=300, msg=False)
    print(f"Status: {pl.LpStatus[status]}")
    if status == pl.LpStatusOptimal:
        optimizer.print_solution()
        analysis = AdvancedAnalysis(optimizer)
        print("\nWorkload:", analysis.workload_distribution())
        print("Completion by priority:", analysis.mission_completion_by_priority())
        print("Sensitivity:")
        for row in optimizer.sensitivity_analysis():
            print(" ", row)


if __name__ == "__main__":
    main()
